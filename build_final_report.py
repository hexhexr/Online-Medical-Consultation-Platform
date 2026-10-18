import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import subprocess
import pypdf
import os
import re

COLOR_BLACK = RGBColor(0, 0, 0)

FONT_BODY_NAME = "Times New Roman"
FONT_BODY_SIZE = Pt(12)      # Strictly 12pt
FONT_H1_SIZE = Pt(16)        # 16pt Bold
FONT_H2_SIZE = Pt(14)        # 14pt Bold
FONT_H3_SIZE = Pt(12)        # 12pt Bold
FONT_CAPTION_SIZE = Pt(10)   # 10pt Bold
FONT_CODE_NAME = "Courier New"
FONT_CODE_SIZE = Pt(9.0)     # 9pt Monospace

FONT_TBL_HDR_SIZE = Pt(10.0) # 10pt Bold
FONT_TBL_CELL_SIZE = Pt(9.5) # 9.5pt Regular

def format_cell(cell, font_size=FONT_TBL_CELL_SIZE, bold=False, line_spacing=1.1, pad_v=20, pad_h=55):
    tcPr = cell._tc.get_or_add_tcPr()
    existing_shd = tcPr.find(qn('w:shd'))
    if existing_shd is not None:
        tcPr.remove(existing_shd)
    
    tcMar = parse_xml(
        r'<w:tcMar %s>'
        r'  <w:top w:w="%s" w:type="dxa"/>'
        r'  <w:bottom w:w="%s" w:type="dxa"/>'
        r'  <w:left w:w="%s" w:type="dxa"/>'
        r'  <w:right w:w="%s" w:type="dxa"/>'
        r'</w:tcMar>' % (nsdecls('w'), pad_v, pad_v, pad_h, pad_h)
    )
    existing_mar = tcPr.find(qn('w:tcMar'))
    if existing_mar is not None:
        tcPr.remove(existing_mar)
    tcPr.append(tcMar)

    for cp in cell.paragraphs:
        cp.paragraph_format.space_before = Pt(0.5)
        cp.paragraph_format.space_after = Pt(0.5)
        cp.paragraph_format.line_spacing = line_spacing
        for cr in cp.runs:
            cr.font.name = FONT_BODY_NAME
            cr.font.color.rgb = COLOR_BLACK
            cr.font.size = font_size
            cr.font.bold = bold

def style_uniform_table(t, font_sz=FONT_TBL_CELL_SIZE, hdr_font_sz=FONT_TBL_HDR_SIZE, line_sp=1.1, pad_v=20, pad_h=55):
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    existing_borders = t._tbl.tblPr.find(qn('w:tblBorders'))
    if existing_borders is not None:
        t._tbl.tblPr.remove(existing_borders)
    t._tbl.tblPr.append(tblBorders)

    for row_idx, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        
        if row_idx == 0:
            if trPr.find(qn('w:tblHeader')) is None:
                trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))

        for cell in row.cells:
            if row_idx == 0:
                format_cell(cell, font_size=hdr_font_sz, bold=True, line_spacing=line_sp, pad_v=pad_v, pad_h=pad_h)
            else:
                format_cell(cell, font_size=font_sz, bold=False, line_spacing=line_sp, pad_v=pad_v, pad_h=pad_h)

def resize_image_in_paragraph(p, new_width_in, new_height_in):
    new_cx = str(int(new_width_in * 914400))
    new_cy = str(int(new_height_in * 914400))
    for r in p.runs:
        if 'w:drawing' in r._r.xml:
            for elem in r._r.iter():
                if 'cx' in elem.attrib and 'cy' in elem.attrib:
                    elem.attrib['cx'] = new_cx
                    elem.attrib['cy'] = new_cy

def set_cell_margins(cell, top=35, bottom=35, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_table_no_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_master_report():
    subprocess.run(['git', 'checkout', 'HEAD', '--', 'FINAL_FYDP_PROJECT_REPORT.docx'], check=True)
    doc = docx.Document('FINAL_FYDP_PROJECT_REPORT.docx')
    body = doc.element.body

    style_normal = doc.styles['Normal']
    style_normal.font.name = FONT_BODY_NAME
    style_normal.font.size = FONT_BODY_SIZE
    style_normal.font.color.rgb = COLOR_BLACK

    for idx, sec in enumerate(doc.sections):
        if idx > 0:
            sec.top_margin = Inches(1.0)
            sec.bottom_margin = Inches(1.0)
            sec.left_margin = Inches(1.0)
            sec.right_margin = Inches(1.0)

    intro_idx = None
    for i, child in enumerate(body):
        if child.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(child, doc)
            if p.text.strip() == '1. Introduction':
                intro_idx = i
                break
    
    for child in list(body[12:intro_idx]):
        body.remove(child)

    intro_elem = body[12]

    def insert_p(text="", space_before=0, space_after=6, line_spacing=1.2, bold=False, italic=False, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_with_next=False, page_break_before=False):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.keep_with_next = keep_with_next
        p.paragraph_format.page_break_before = page_break_before
        if text:
            r = p.add_run(text)
            r.font.name = FONT_BODY_NAME
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = COLOR_BLACK
        intro_elem.addprevious(p._p)
        return p

    def insert_heading_1(text, page_break_before=True):
        return insert_p(text, space_before=14, space_after=6, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.LEFT, keep_with_next=True, page_break_before=page_break_before)

    def insert_toc_line(title, page_num, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.0)
        p.paragraph_format.space_after = Pt(1.0)
        p.paragraph_format.line_spacing = 1.12
        indent = 0.0 if level == 1 else 0.25
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        
        r1 = p.add_run(title)
        r1.font.name = FONT_BODY_NAME
        r1.font.size = FONT_BODY_SIZE
        r1.font.bold = (level == 1)
        r1.font.color.rgb = COLOR_BLACK
        
        r2 = p.add_run(f"\t{page_num}")
        r2.font.name = FONT_BODY_NAME
        r2.font.size = FONT_BODY_SIZE
        r2.font.bold = (level == 1)
        r2.font.color.rgb = COLOR_BLACK
        
        intro_elem.addprevious(p._p)
        return p

    def insert_list_line(title, page_num):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        
        r1 = p.add_run(title)
        r1.font.name = FONT_BODY_NAME
        r1.font.size = FONT_BODY_SIZE
        r1.font.color.rgb = COLOR_BLACK
        
        r2 = p.add_run(f"\t{page_num}")
        r2.font.name = FONT_BODY_NAME
        r2.font.size = FONT_BODY_SIZE
        r2.font.color.rgb = COLOR_BLACK
        
        intro_elem.addprevious(p._p)
        return p

    # --- FRONT MATTER: AUTHOR'S DECLARATION (Single Page) ---
    p_dec_head = doc.add_paragraph()
    p_dec_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dec_head.paragraph_format.space_before = Pt(8)
    p_dec_head.paragraph_format.space_after = Pt(12)
    p_dec_head.paragraph_format.page_break_before = True
    p_dec_head.paragraph_format.keep_with_next = True
    r_dec_head = p_dec_head.add_run("AUTHOR’S DECLARATION")
    r_dec_head.font.name = FONT_BODY_NAME
    r_dec_head.font.size = Pt(14)
    r_dec_head.font.bold = True
    r_dec_head.font.color.rgb = COLOR_BLACK
    intro_elem.addprevious(p_dec_head._p)

    p_d1 = doc.add_paragraph()
    p_d1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_d1.paragraph_format.line_spacing = 1.2
    p_d1.paragraph_format.space_before = Pt(0)
    p_d1.paragraph_format.space_after = Pt(6)
    r_d1 = p_d1.add_run(
        "I hereby declare that I am the sole author of this thesis. This is a true copy of the thesis, "
        "including any required final revisions, as accepted by my examiners. It is further declared "
        "that this thesis or any part thereof has not been previously submitted for any degree or diploma "
        "in any university or other educational institution."
    )
    r_d1.font.name = FONT_BODY_NAME
    r_d1.font.size = FONT_BODY_SIZE
    r_d1.font.color.rgb = COLOR_BLACK
    intro_elem.addprevious(p_d1._p)

    p_d2 = doc.add_paragraph()
    p_d2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_d2.paragraph_format.line_spacing = 1.2
    p_d2.paragraph_format.space_before = Pt(0)
    p_d2.paragraph_format.space_after = Pt(10)
    r_d2 = p_d2.add_run("I understand that my thesis may be made electronically available to the public.")
    r_d2.font.name = FONT_BODY_NAME
    r_d2.font.size = FONT_BODY_SIZE
    r_d2.font.color.rgb = COLOR_BLACK
    intro_elem.addprevious(p_d2._p)

    decl_data = [
        ("Student Name", "Muhammad Danial , Muhammad Bilal , Mudassar Javed"),
        ("Registration No.", ["BSCSW-2.5Y-021R24-4 , BSCSW-2.5Y-021R24-1 ,", "BSCSW-2.5Y-021R24-5"]),
        ("Contact No.", "03097705525, 03036556495, 03014521025"),
        ("Email Address", "digdaguu@gmail.com, bilalsadiq612@gmail.com, mudassar21025@gmail.com"),
        ("Discipline", "Computer Science"),
        ("Year of Study", "(2024-2026)"),
        ("Area of Study", "BSCS"),
        ("Proposed Project Title", "Online Medical Consultation Platform"),
    ]
    t_decl = doc.add_table(rows=len(decl_data), cols=2)
    t_decl.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_decl.autofit = False
    t_decl.allow_autofit = False
    set_table_borders(t_decl, color="000000", sz="4", val="single")

    col_widths = [Inches(1.8), Inches(4.7)]
    for row_idx, (label, val) in enumerate(decl_data):
        row = t_decl.rows[row_idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        
        c0 = row.cells[0]
        c0.width = col_widths[0]
        set_cell_margins(c0, top=35, bottom=35, left=70, right=70)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.line_spacing = 1.15
        r0 = p0.add_run(label)
        r0.font.name = FONT_BODY_NAME
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_BLACK

        c1 = row.cells[1]
        c1.width = col_widths[1]
        set_cell_margins(c1, top=35, bottom=35, left=70, right=70)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.15
        if isinstance(val, list):
            for v_idx, v_item in enumerate(val):
                if v_idx > 0:
                    p_sub = c1.add_paragraph()
                    p_sub.paragraph_format.space_before = Pt(1)
                    p_sub.paragraph_format.space_after = Pt(0)
                    p_sub.paragraph_format.line_spacing = 1.15
                    r_sub = p_sub.add_run(v_item)
                    r_sub.font.name = FONT_BODY_NAME
                    r_sub.font.size = Pt(10.5)
                    r_sub.font.color.rgb = COLOR_BLACK
                else:
                    r1 = p1.add_run(v_item)
                    r1.font.name = FONT_BODY_NAME
                    r1.font.size = Pt(10.5)
                    r1.font.color.rgb = COLOR_BLACK
        else:
            r1 = p1.add_run(val)
            r1.font.name = FONT_BODY_NAME
            r1.font.size = Pt(10.5)
            r1.font.color.rgb = COLOR_BLACK

    intro_elem.addprevious(t_decl._tbl)

    p_sig_sp = doc.add_paragraph()
    p_sig_sp.paragraph_format.space_before = Pt(18)
    p_sig_sp.paragraph_format.space_after = Pt(0)
    p_sig_sp.paragraph_format.line_spacing = 1.0
    intro_elem.addprevious(p_sig_sp._p)

    t_sig = doc.add_table(rows=1, cols=2)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sig.autofit = False
    t_sig.allow_autofit = False
    set_table_no_borders(t_sig)

    sig_widths = [Inches(3.25), Inches(3.25)]
    row_sig = t_sig.rows[0]
    cell_s = row_sig.cells[0]
    cell_s.width = sig_widths[0]
    set_cell_margins(cell_s, top=0, bottom=0, left=0, right=10)
    ps = cell_s.paragraphs[0]
    ps.paragraph_format.space_before = Pt(0)
    ps.paragraph_format.space_after = Pt(3)
    rs_lbl = ps.add_run("Student Signature: ")
    rs_lbl.font.name = FONT_BODY_NAME
    rs_lbl.font.size = Pt(11)
    rs_lbl.font.bold = True
    rs_lbl.font.color.rgb = COLOR_BLACK
    rs_line = ps.add_run("__________________")
    rs_line.font.name = FONT_BODY_NAME
    rs_line.font.size = Pt(11)
    rs_line.font.color.rgb = COLOR_BLACK
    for s_name in ["Muhammad Danial", "Muhammad Bilal", "Mudassar Javed"]:
        ps_n = cell_s.add_paragraph()
        ps_n.paragraph_format.space_before = Pt(2)
        ps_n.paragraph_format.space_after = Pt(0)
        rn = ps_n.add_run(s_name)
        rn.font.name = FONT_BODY_NAME
        rn.font.size = Pt(10.5)
        rn.font.color.rgb = COLOR_BLACK

    cell_sup = row_sig.cells[1]
    cell_sup.width = sig_widths[1]
    set_cell_margins(cell_sup, top=0, bottom=0, left=10, right=0)
    psup = cell_sup.paragraphs[0]
    psup.paragraph_format.space_before = Pt(0)
    psup.paragraph_format.space_after = Pt(3)
    rsup_lbl = psup.add_run("Supervisor Signature: ")
    rsup_lbl.font.name = FONT_BODY_NAME
    rsup_lbl.font.size = Pt(11)
    rsup_lbl.font.bold = True
    rsup_lbl.font.color.rgb = COLOR_BLACK
    rsup_line = psup.add_run("__________________")
    rsup_line.font.name = FONT_BODY_NAME
    rsup_line.font.size = Pt(11)
    rsup_line.font.color.rgb = COLOR_BLACK
    psup_n = cell_sup.add_paragraph()
    psup_n.paragraph_format.space_before = Pt(2)
    psup_n.paragraph_format.space_after = Pt(0)
    psup_n.paragraph_format.left_indent = Inches(1.2)
    rsup_name = psup_n.add_run("Majid Khawar")
    rsup_name.font.name = FONT_BODY_NAME
    rsup_name.font.size = Pt(11)
    rsup_name.font.bold = True
    rsup_name.font.color.rgb = COLOR_BLACK

    intro_elem.addprevious(t_sig._tbl)

    # --- FRONT MATTER: ABSTRACT (Single cohesive paragraph, 12pt) ---
    insert_heading_1("Abstract", page_break_before=True)
    abstract_content = (
        "Access to specialized healthcare in developing countries such as Pakistan remains severely constrained by "
        "geographic barriers, acute physician shortages (approximately 1 doctor per 1,000 citizens according to World Health "
        "Organization metrics), and prohibitive consultation expenses. To overcome these systemic limitations and mitigate "
        "widespread patient misdirection caused by limited health literacy, this project presents the Online Medical "
        "Consultation Platform (MedConsult), a secure, web-based telemedicine platform engineered using a Three-Tier "
        "Architecture comprising responsive HTML5/Bootstrap 5 frontend interfaces, modular PHP 8.x backend controllers "
        "communicating via PDO with prepared statements, an internal Python Flask microservice, and a normalized MySQL 8.0 "
        "relational database. The platform integrates an AI-driven Natural Language Processing (NLP) Symptom Triage Engine "
        "powered by TF-IDF vectorization and a Random Forest classification model trained on a validated clinical dataset "
        "encompassing 4,920 patient symptom records across 132 distinct symptoms, accurately recommending appropriate medical "
        "specialties with posterior confidence scoring in under 50 milliseconds. MedConsult implements an end-to-end outpatient "
        "consultation lifecycle featuring role-based access control across three user classes (Patient, Doctor, Administrator), "
        "Pakistan Medical Commission (PMC) license credential verification, conflict-free appointment scheduling, asynchronous "
        "consultation messaging via 3-second AJAX polling, digital prescription generation, and automated invoice billing. "
        "Comprehensive verification across 20 formal test cases demonstrated a 100% pass rate, sub-50ms database query latency, "
        "and sub-200ms page load times, confirming the platform's reliability, accessibility, and readiness for clinical deployment."
    )
    insert_p(abstract_content, space_before=0, space_after=12, line_spacing=1.35, size=12)

    # --- FRONT MATTER: ACKNOWLEDGEMENTS (12pt) ---
    insert_heading_1("Acknowledgements", page_break_before=True)
    insert_p("First and foremost, we express our profound gratitude to Almighty Allah for bestowing upon us the intellect, perseverance, and strength required to successfully design, develop, and complete this Final Year Design Project.", space_after=8, size=12)
    insert_p("We wish to express our deepest and most sincere appreciation to our project supervisor, Majid Khawar, whose guidance, technical mentorship, and constructive feedback were instrumental in shaping this platform from requirement formulation to production deployment. His profound insights into software architecture and practical engineering significantly elevated the maturity and rigor of our work.", space_after=8, size=12)
    insert_p("We also extend our heartfelt thanks to the faculty members and laboratory staff of the Department of Computer Science & Information Technology at the University of Southern Punjab, Multan, for providing an academically stimulating environment, computational resources, and constructive reviews during our project evaluations.", space_after=8, size=12)
    insert_p("Finally, we are eternally indebted to our beloved parents and families for their endless sacrifices, prayers, and unconditional encouragement throughout our academic journey. We also thank our classmates for their collaborative spirit and valuable suggestions.", space_after=12, size=12)

    # --- FRONT MATTER: TABLE OF CONTENTS (12pt, perfectly fitted on Page iv) ---
    insert_heading_1("Table of Contents", page_break_before=True)
    toc_data = [
        ("1. Introduction", "1", 1),
        ("1.1 Background", "1", 2),
        ("1.2 Problem Statement", "1", 2),
        ("1.3 Objectives", "2", 2),
        ("1.4 Scope", "2", 2),
        ("1.5 Significance", "3", 2),
        ("1.6 Organization of the Report", "3", 2),
        ("2. Literature Review", "4", 1),
        ("2.1 Review of Related Work", "4", 2),
        ("2.2 Gap Analysis", "5", 2),
        ("2.3 Summary of Theoretical Foundations", "6", 2),
        ("3. Methodology", "7", 1),
        ("3.1 Project Design", "7", 2),
        ("3.2 Technology Stack", "7", 2),
        ("3.3 Implementation", "8", 2),
        ("3.4 Testing", "9", 2),
        ("4. System Design", "11", 1),
        ("4.1 Design Overview", "11", 2),
        ("4.2 Architecture", "11", 2),
        ("4.3 Data Design", "15", 2),
        ("4.3.1 Relational Data Dictionaries", "16", 2),
        ("5. Results and Discussion", "20", 1),
        ("5.1 Results", "20", 2),
        ("5.1.3 Operational User Interface Walkthrough", "21", 2),
        ("5.2 Discussion", "26", 2),
        ("6. Conclusion", "27", 1),
        ("6.1 Summary", "27", 2),
        ("6.2 Contributions", "27", 2),
        ("6.3 Future Work", "28", 2),
        ("References", "29", 1),
        ("Appendices", "30", 1),
        ("Appendix A: Production Cloud Deployment & Infrastructure Guide", "30", 2),
        ("Appendix B: System User Guide & Operating Manual", "30", 2),
        ("Appendix C: Core System Source Code Listings", "31", 2),
        ("Glossary", "32", 1),
    ]
    for title, pg, lvl in toc_data:
        insert_toc_line(title, pg, lvl)

    # --- FRONT MATTER: LIST OF FIGURES (12pt, perfectly fitted on Page v) ---
    insert_heading_1("List of Figures", page_break_before=True)
    figures_data = [
        ("Figure 4.1: Three-Tier System Architecture Diagram", "11"),
        ("Figure 4.2: Context Diagram (DFD Level 0)", "12"),
        ("Figure 4.3: Complete System Use Case Diagram", "12"),
        ("Figure 4.4: Sequence Diagram — AI Symptom Triage Workflow", "13"),
        ("Figure 4.5: Sequence Diagram — Appointment Booking Lifecycle", "13"),
        ("Figure 4.6: Sequence Diagram — Real-Time Consultation Messaging", "14"),
        ("Figure 4.7: System Activity Workflow Diagram", "14"),
        ("Figure 4.8: Entity Relationship Diagram (ERD)", "15"),
        ("Figure 5.1: Public Landing Page and Service Overview", "21"),
        ("Figure 5.2: Searchable Doctor Directory with Specialty Filters", "21"),
        ("Figure 5.3: Unified Authentication Portal (Patient / Doctor / Admin)", "22"),
        ("Figure 5.4: Patient Dashboard with Recent Activity", "22"),
        ("Figure 5.5: AI NLP Symptom Triage Interface with Specialist Recommendation", "23"),
        ("Figure 5.6: Patient Appointment Management and Real-Time Chat Modal", "23"),
        ("Figure 5.7: Patient Medical History and Consultation Records Archive", "24"),
        ("Figure 5.8: Doctor Clinical Dashboard and Statistics Overview", "24"),
        ("Figure 5.9: Doctor Appointment Review and Digital Prescription Issuance", "25"),
        ("Figure 5.10: Administrator Control Panel and Doctor PMC License Audit", "25"),
    ]
    for title, pg in figures_data:
        insert_list_line(title, pg)

    # --- FRONT MATTER: LIST OF TABLES (12pt, perfectly fitted on Page vi) ---
    insert_heading_1("List of Tables", page_break_before=True)
    tables_data = [
        ("Table 2.1: Comparative Feature Matrix of Telemedicine Platforms", "5"),
        ("Table 3.1: Agile Scrum Development Sprint Schedule and Deliverables", "7"),
        ("Table 3.2: Technology Stack Specifications and Architectural Roles", "7"),
        ("Table 3.3: Formal Verification Test Cases Matrix (TC-01 to TC-20)", "9"),
        ("Table 4.1: Data Dictionary — Users Entity (users)", "16"),
        ("Table 4.2: Data Dictionary — Appointments Entity (appointments)", "17"),
        ("Table 4.3: Data Dictionary — Messages Entity (messages)", "17"),
        ("Table 4.4: Data Dictionary — Prescriptions Entity (prescriptions)", "18"),
        ("Table 4.5: Data Dictionary — Invoices Entity (invoices)", "18"),
        ("Table 4.6: Data Dictionary — Audit Log Entity (audit_log)", "19"),
        ("Table 4.7: Data Dictionary — Symptoms Reference Entity (symptoms)", "19"),
        ("Table 4.8: Data Dictionary — Contact Inquiries Entity (contact_messages)", "19"),
        ("Table 5.1: AI Symptom Triage Model Evaluation Metrics by Clinical Specialty", "20"),
        ("Table 5.2: System Performance and Latency Benchmarks", "20"),
    ]
    for title, pg in tables_data:
        insert_list_line(title, pg)

    # Section break between Front Matter and Main Body
    sec_break_p = doc.add_paragraph()
    sec_break_pPr = sec_break_p._p.get_or_add_pPr()
    sec_break_sectPr = parse_xml(
        r'<w:sectPr %s>'
        r'  <w:type w:val="nextPage"/>'
        r'  <w:pgNumType w:fmt="decimal" w:start="1"/>'
        r'</w:sectPr>' % nsdecls('w')
    )
    sec_break_pPr.append(sec_break_sectPr)
    intro_elem.addprevious(sec_break_p._p)

    # DEDUPLICATION: REMOVE ONLY KNOWN DUPLICATES
    all_paras = list(doc.paragraphs)
    to_remove = []

    count_1_6 = 0
    in_dup_1_6 = False
    count_2_3 = 0
    count_c1 = 0
    count_c2 = 0

    for p in all_paras:
        txt = p.text.strip()
        has_draw = 'w:drawing' in p._p.xml
        if has_draw:
            continue

        if not txt:
            if p._p.pPr is None or p._p.pPr.sectPr is None:
                to_remove.append(p._p)
            continue

        if txt == "1.6 Organization of the Report":
            count_1_6 += 1
            if count_1_6 > 1:
                in_dup_1_6 = True
                to_remove.append(p._p)
                continue
        
        if in_dup_1_6:
            if txt == "2. Literature Review":
                in_dup_1_6 = False
            else:
                to_remove.append(p._p)
                continue

        if txt == "2.3 Summary of Theoretical Foundations":
            count_2_3 += 1
            if count_2_3 > 1:
                to_remove.append(p._p)
                continue

        if count_2_3 > 1 and txt.startswith("The literature review and gap analysis establish clear architectural"):
            to_remove.append(p._p)
            continue

        if txt.startswith("Listing C.1:"):
            count_c1 += 1
            if count_c1 > 1:
                to_remove.append(p._p)
                continue

        if count_c1 > 1 and txt.startswith("@app.route('/predict'"):
            to_remove.append(p._p)
            continue

        if txt.startswith("Listing C.2:"):
            count_c2 += 1
            if count_c2 > 1:
                to_remove.append(p._p)
                continue

        if count_c2 > 1 and txt.startswith("function get_db_connection()"):
            to_remove.append(p._p)
            continue

    sovereign_count = 0
    for p in all_paras:
        if p.text.strip().startswith("• Lack of Sovereign Open-Source Infrastructure:"):
            sovereign_count += 1
            if sovereign_count > 1:
                to_remove.append(p._p)

    for node in to_remove:
        if node.getparent() is not None:
            node.getparent().remove(node)

    # ADD ENRICHMENT TO CHAPTER 4 DATABASE AND CHAPTER 6 CONCLUSION
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == "4.3.1 Relational Data Dictionaries":
            p_extra = doc.add_paragraph()
            p_extra.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_extra.paragraph_format.line_spacing = 1.15
            p_extra.paragraph_format.space_before = Pt(0)
            p_extra.paragraph_format.space_after = Pt(3)
            r_ex = p_extra.add_run(
                "To guarantee high transaction throughput and sub-50ms query execution under concurrent user load, "
                "composite B-tree indexes were engineered across frequently filtered columns including "
                "users(role, is_verified, specialty), appointments(doctor_id, preferred_date, status), and "
                "messages(appointment_id, created_at). Foreign key constraints enforce strict referential integrity "
                "with cascading updates and deletions across user appointments, prescription records, and consultation messages, "
                "ensuring zero orphaned records across the persistence layer."
            )
            r_ex.font.name = FONT_BODY_NAME
            r_ex.font.size = FONT_BODY_SIZE
            r_ex.font.color.rgb = COLOR_BLACK
            p._p.addprevious(p_extra._p)
            break

    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == "References":
            p_future = doc.add_paragraph()
            p_future.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_future.paragraph_format.line_spacing = 1.15
            p_future.paragraph_format.space_before = Pt(4)
            p_future.paragraph_format.space_after = Pt(4)
            r_fut = p_future.add_run(
                "In synthesis, the future work roadmap positions MedConsult not merely as a standalone academic design project, "
                "but as an extensible, enterprise-grade telemedicine foundation tailored specifically to the structural and "
                "socio-economic realities of Pakistan's healthcare ecosystem. By advancing from specialized outpatient routing "
                "toward multilingual vernacular support, peer-to-peer visual consultation, cross-platform mobile accessibility, "
                "and hospital EHR federation, MedConsult provides a sustainable blueprint for national digital health equity "
                "and sovereign clinical innovation."
            )
            r_fut.font.name = FONT_BODY_NAME
            r_fut.font.size = FONT_BODY_SIZE
            r_fut.font.color.rgb = COLOR_BLACK
            p._p.addprevious(p_future._p)
            break

    # RENAME "14. Glossary" -> "Glossary"
    for p in doc.paragraphs:
        if p.text.strip() == "14. Glossary":
            p.text = "Glossary"

    # FORMAT ALL ELEMENTS IN MAIN BODY (SECTION 2)
    in_section_2 = False

    for p in doc.paragraphs:
        txt = p.text.strip()
        has_draw = 'w:drawing' in p._p.xml
        
        if txt == "1. Introduction":
            in_section_2 = True
            
        if not in_section_2:
            continue

        if has_draw:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.page_break_before = False
            continue

        if not txt:
            continue

        # Monospace Code: 9pt Courier New
        if any(r.font.name == FONT_CODE_NAME for r in p.runs) or txt.startswith("@app.route") or txt.startswith("function get_db_connection"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.page_break_before = False
            for r in p.runs:
                r.font.name = FONT_CODE_NAME
                r.font.size = FONT_CODE_SIZE
                r.font.color.rgb = COLOR_BLACK
            continue

        # Captions: 10pt Bold (Italic for Figures)
        if txt.startswith("Figure ") and ":" in txt:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.page_break_before = False
            for r in p.runs:
                r.font.name = FONT_BODY_NAME
                r.font.bold = True
                r.font.italic = True
                r.font.size = FONT_CAPTION_SIZE
                r.font.color.rgb = COLOR_BLACK
            continue

        if txt.startswith("Table ") and ":" in txt:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.page_break_before = False
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = FONT_BODY_NAME
                r.font.bold = True
                r.font.size = FONT_CAPTION_SIZE
                r.font.color.rgb = COLOR_BLACK
            continue

        if txt.startswith("Listing ") and ":" in txt:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.page_break_before = False
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = FONT_BODY_NAME
                r.font.bold = True
                r.font.size = FONT_CAPTION_SIZE
                r.font.color.rgb = COLOR_BLACK
            continue

        # Figure intro sentences: keep with next
        if txt.startswith("Figure ") and any(w in txt for w in ['illustrates', 'presents', 'depicts', 'shows', 'displays', 'demonstrates']):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(2.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.page_break_before = False
            for r in p.runs:
                r.font.name = FONT_BODY_NAME
                r.font.size = FONT_BODY_SIZE
                r.font.color.rgb = COLOR_BLACK
            continue

        # Heading 1 (Chapters & Main Headings): 16pt Bold on a FRESH page
        if txt in [
            "1. Introduction", "2. Literature Review", "3. Methodology",
            "4. System Design", "5. Results and Discussion", "6. Conclusion",
            "References", "Appendices", "Glossary"
        ]:
            p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = FONT_BODY_NAME
                r.font.size = FONT_H1_SIZE
                r.font.bold = True
                r.font.color.rgb = COLOR_BLACK
            continue

        # Heading 2 (Major Sections & Appendices): 14pt Bold
        if any(txt.startswith(f"{c}.{s} ") for c in range(1, 7) for s in range(1, 10)) or txt.startswith("Appendix "):
            p.paragraph_format.page_break_before = txt.startswith("6.3 Future Work")
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = FONT_BODY_NAME
                r.font.size = FONT_H2_SIZE
                r.font.bold = True
                r.font.color.rgb = COLOR_BLACK
            continue

        # Heading 3 (Subsections): 12pt Bold
        if any(txt.startswith(f"{c}.{s}.{ss} ") for c in range(1, 7) for s in range(1, 10) for ss in range(1, 10)):
            p.paragraph_format.page_break_before = (txt.startswith("4.3.1 Relational Data Dictionaries"))
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = FONT_BODY_NAME
                r.font.size = FONT_H3_SIZE
                r.font.bold = True
                r.font.color.rgb = COLOR_BLACK
            continue

        # Glossary entries: 12pt, compact spacing
        if in_section_2 and any(txt.startswith(g) for g in ["API (", "AJAX (", "Backend:", "BCrypt:", "Database Schema:", "DFD (", "ERD (", "Frontend:", "GUI (", "IP Address:", "Machine Learning (", "NLP (", "PDO (", "PMC (", "RBAC (", "TF-IDF (", "Triage:", "XAMPP:"]):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.10
            p.paragraph_format.space_before = Pt(0.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.page_break_before = False
            for r in p.runs:
                r.font.name = FONT_BODY_NAME
                r.font.size = FONT_BODY_SIZE
                r.font.color.rgb = COLOR_BLACK
            continue

        # Normal Body Paragraphs: Strictly 12pt Times New Roman, line spacing 1.15, space after 2.5pt
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.page_break_before = False
        for r in p.runs:
            r.font.name = FONT_BODY_NAME
            r.font.size = FONT_BODY_SIZE
            r.font.color.rgb = COLOR_BLACK

    # CHAPTER 4 & 5 SCREENSHOT & DIAGRAM PAIRING
    in_section_2 = False
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == "1. Introduction":
            in_section_2 = True
        if not in_section_2:
            continue

        if "Figure 4.2 presents the Context Diagram" in txt or \
           "Figure 4.4 illustrates the sequence of messages" in txt or \
           "Figure 4.6 illustrates Consultation Messaging" in txt or \
           "4.3 Data Design" in txt or \
           "Table 4.2: Data Dictionary" in txt or \
           "Table 4.4: Data Dictionary" in txt or \
           "Table 4.6: Data Dictionary" in txt:
            p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True

        if "5.1.3 Operational User Interface Walkthrough" in txt or \
           "Figure 5.3 shows the unified authentication portal" in txt or \
           "Figure 5.5 demonstrates the AI NLP Symptom Triage" in txt or \
           "Figure 5.7 depicts the patient medical records archive" in txt or \
           "Figure 5.9 demonstrates the doctor consultation review" in txt or \
           "5.2 Discussion" in txt:
            p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True

    # RESIZE IMAGES
    for i, p in enumerate(doc.paragraphs):
        has_img = 'w:drawing' in p._p.xml and i >= 12
        if not has_img:
            continue
        prev_txt = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        next_txt = doc.paragraphs[i+1].text.strip() if i < len(doc.paragraphs)-1 else ""

        if "4.1" in prev_txt or "4.1" in next_txt:
            resize_image_in_paragraph(p, 4.8, 2.7)
        elif "4.2" in prev_txt or "4.2" in next_txt or "4.3" in prev_txt or "4.3" in next_txt:
            resize_image_in_paragraph(p, 4.4, 2.3)
        elif "4.4" in prev_txt or "4.4" in next_txt or "4.5" in prev_txt or "4.5" in next_txt:
            resize_image_in_paragraph(p, 4.4, 2.3)
        elif "4.6" in prev_txt or "4.6" in next_txt or "4.7" in prev_txt or "4.7" in next_txt:
            resize_image_in_paragraph(p, 4.4, 2.3)
        elif "4.8" in prev_txt or "4.8" in next_txt:
            resize_image_in_paragraph(p, 4.8, 2.6)
        elif any(f"5.{k}" in prev_txt or f"5.{k}" in next_txt for k in range(1, 11)):
            resize_image_in_paragraph(p, 4.6, 2.25)

    # UNIFORM TABLE FORMATTING (Table 2.1 to Table 5.2):
    for t_idx, t in enumerate(doc.tables):
        if t_idx == 0 and len(t.rows) == 8:
            continue
        if t_idx == 1 and len(t.rows) == 1 and len(t.columns) == 2:
            continue
        style_uniform_table(t, font_sz=FONT_TBL_CELL_SIZE, hdr_font_sz=FONT_TBL_HDR_SIZE, line_sp=1.1, pad_v=20, pad_h=55)

    # HEADERS & FOOTERS:
    for s_idx, s in enumerate(doc.sections):
        s.header.is_linked_to_previous = False
        for hp in s.header.paragraphs:
            hp.text = ''
            pPr = hp._p.get_or_add_pPr()
            pBdr = pPr.find(qn('w:pBdr'))
            if pBdr is not None: pPr.remove(pBdr)

        if s_idx == 0:
            s.different_first_page_header_footer = True
            for fp in s.first_page_footer.paragraphs: fp.text = ''
            for fp in s.first_page_header.paragraphs: fp.text = ''
            for p in s.footer.paragraphs: p.text = ''
        else:
            s.different_first_page_header_footer = False
            s.footer.is_linked_to_previous = False
            footer = s.footer
            fp = footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fp.text = ''
            pPr = fp._p.get_or_add_pPr()
            pBdr = pPr.find(qn('w:pBdr'))
            if pBdr is not None: pPr.remove(pBdr)
            r = fp.add_run()
            r.font.name = FONT_BODY_NAME
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_BLACK
            fp._p.append(parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w')))

    if len(doc.sections) >= 2:
        s1 = doc.sections[1]
        s1_sectPr = s1._sectPr
        existing_pgNum = s1_sectPr.find(qn('w:pgNumType'))
        if existing_pgNum is not None: s1_sectPr.remove(existing_pgNum)
        s1_sectPr.append(parse_xml(r'<w:pgNumType %s w:fmt="lowerRoman" w:start="1"/>' % nsdecls('w')))

    if len(doc.sections) >= 3:
        s2 = doc.sections[2]
        s2_sectPr = s2._sectPr
        existing_pgNum2 = s2_sectPr.find(qn('w:pgNumType'))
        if existing_pgNum2 is not None: s2_sectPr.remove(existing_pgNum2)
        s2_sectPr.append(parse_xml(r'<w:pgNumType %s w:fmt="decimal" w:start="1"/>' % nsdecls('w')))

    output_docx = 'FINAL_FYDP_PROJECT_REPORT.docx'
    doc.save(output_docx)
    print(f"Successfully generated {output_docx}")

if __name__ == '__main__':
    build_master_report()
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', 'FINAL_FYDP_PROJECT_REPORT.docx'], check=True)
    print("Successfully converted FINAL_FYDP_PROJECT_REPORT.pdf")
